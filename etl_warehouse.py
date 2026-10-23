"""ETL 2 แหล่ง -> Data Warehouse (MySQL)

แหล่งที่ 1  Excel : ใบแจ้งซ่อม Google Form (sheet "การตอบแบบฟอร์ม 1/3/4") -> fact_general_maintenance
แหล่งที่ 2  API   : ระบบครุภัณฑ์ CPEAssets (assets = API จริง, endpoint อื่น = mock_api/*.json)
                    -> dim_asset ฯลฯ, fact_asset_inventory (snapshot), fact_asset_maintenance,
                       fact_borrow_transaction, fact_general_maintenance
ผลลัพธ์ (<out-dir>): warehouse_load.sql, warehouse/*.csv, raw/*.json (API ที่ดึงมา),
                    history/assets_YYYYMMDD.json (snapshot ครุภัณฑ์ทุกรอบ ห้ามลบ),
                    repair_requests_clean.xlsx/.csv (Excel ที่ clean แล้ว + cleaning_log + rejected_place)
"""
import argparse
import json
import os
import re
import sys
import urllib.request
from dataclasses import dataclass, field
from datetime import datetime, timedelta

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TARGET = ["แจ้งซ่อมเมื่อ", "สังกัด", "สถานภาพผู้แจ้งซ่อม", "ประเภทงานซ่อมบำรุง",
          "อาคาร", "ชั้น", "ห้อง", "ชื่อพื้นที่", "รายละเอียดของปัญหา"]
DROP_INVALID_PLACE = True   # False = เก็บแถวไว้แต่เว้นสถานที่ว่าง
RAW = ["ts", "name", "dept", "status", "phone", "type", "place", "detail", "reported", "extra", "img"]

# ---------- mapping ----------
TYPE_MAP = {
    "ทดสอบระบบ": None, "ทดสอบระบบแจ้งซ่อม": None,
    "งานคอมพิวเตอร์": "งานคอมพิวเตอร์",
    "คอมพิวเตอร์ประจำห้องเรียน": "งานคอมพิวเตอร์",
    "หลอดไฟ/สตาร์ทเตอร์เสีย": "งานระบบไฟฟ้า/แสงสว่าง",
    "ห้องน้ำ ชาย และ หญิง": "งานสุขภัณฑ์",
    "กล้องวงจรปิด": "งานกล้องวงจรปิด",
    "ระบบกล้องวงจรปิด": "งานกล้องวงจรปิด",
    "หน้าต่าง,กระจก": "งานอาคาร/สถานที่",
    "ถังดับเพลิงแรงดันต่ำ": "งานอาคาร/สถานที่",
    "โปรเจคเตอร์": "งานโสตทัศนูปกรณ์",
    "รถตู้คณะ ทะเบียน นค3495": "งานอื่นๆ",
    "เครื่องถ่ายเอกสาร": "งานอื่นๆ",
    "Scan นิ้วมือ": "งานอื่นๆ",
}
STD_TYPES = {"งานระบบปรับอากาศ", "งานระบบไฟฟ้า/แสงสว่าง", "งานโสตทัศนูปกรณ์", "งานอาคาร/สถานที่",
             "งานระบบโทรศัพท์", "งานระบบเครือข่าย/อินเตอร์เน็ต", "งานประปา", "งานสุขภัณฑ์"}
STATUS_STD = {"เจ้าหน้าที่", "อาจารย์", "นิสิต", "ช่างเทคนิค", "แม่บ้าน"}
STATUS_MAP = {
    "เจ้าหน้าที่ห้องโลจิสติกส์": "เจ้าหน้าที่",
    "รองคณบดีฝ่ายบริหาร": "อาจารย์",
}
NULL_EXTRA = {"-", "--", ".", "ไม่มี", "ไม่มีค่ะ", "ไม่มีครับ", "n/a", "na", "none", "null"}

DATE_FALLBACK_AFTER = timedelta(days=7)    # แจ้งซ่อมเมื่อ เก่ากว่า timestamp เกินนี้ = พิมพ์ผิด
DATE_FUTURE_TOLERANCE = timedelta(minutes=5)  # ห้ามอยู่หลัง timestamp เกินนี้

# ---------- แยกสถานที่ (อาคาร / ชั้น / ห้อง / ชื่อพื้นที่) ----------
# รหัสห้อง (E9203): E = วิศวกรรมศาสตร์ (ทุกห้องอยู่ใน E อยู่แล้ว จึงตัดทิ้ง)
#   หลักที่ 1 = อาคาร (1-9), หลักที่ 2 = ชั้น, 2 หลักท้าย = เลขห้อง  ->  E9203 = อาคาร 9 ชั้น 2 ห้อง 03
# ใช้ไม่ได้ (valid=False): ไม่มีอาคาร/ห้อง/ชื่อพื้นที่เลย, เป็นอาการเสียไม่ใช่สถานที่,
#   ชื่อทั่วไปที่ไม่มีอาคาร, หรือ อาคาร/ชั้น ที่พิมพ์ขัดกับรหัสห้อง
VALID_BUILDINGS = set(range(1, 10))

# อาคารที่เรียกด้วยชื่อ -> ชื่อมาตรฐาน (สโมสรนิสิต = อาคารกิจการนิสิต: assumption, ห้อง 2/4 = 204)
NAMED_BUILDINGS = [
    (r"อาคาร\s*กิจการนิสิต(?:ใหม่)?", "อาคารกิจการนิสิต"),
    (r"อาคาร\s*สโมสรนิสิต(?:ใหม่)?", "อาคารกิจการนิสิต"),
    (r"อาคาร\s*สโม(?!สร)", "อาคารกิจการนิสิต"),
    (r"(?<!ห้อง)(?<!หน้าห้อง)สโมสรนิสิตใหม่", "อาคารกิจการนิสิต"),
    (r"อาคาร\s*[เอ]นกประสงค์|อาคารอเนกประสงค์", "อาคารเอนกประสงค์"),
]

# ชื่อพื้นที่ (ลำดับสำคัญ: อันแรกที่ match ชนะ) -> ชื่อมาตรฐาน
AREAS = [
    (r"งานบริการ", "ห้องงานบริการการศึกษาและพัฒนานิสิต"),
    (r"โลจิสติกส์", "สำนักประสานงานโลจิสติกส์"),
    (r"สำนักงานเลขา|office", "สำนักงานเลขานุการ"),
    (r"server|เซิร์ฟเวอร์|เซิฟเวอร์", "ห้อง Server"),
    (r"ห้องน้ำ\s*ชาย\s*(?:และ|,)?\s*(?:ห้องน้ำ)?\s*หญิง", "ห้องน้ำชาย, ห้องน้ำหญิง"),
    (r"ห้องน้ำ\s*หญิง", "ห้องน้ำหญิง"),
    (r"ห้องน้ำ\s*ชาย", "ห้องน้ำชาย"),
    (r"ห้องน้ำ", "ห้องน้ำ"),
    (r"ห้องธุรการ|ธุรการ", "ห้องธุรการ"),
    (r"ห้องพักรองคณบดี", "ห้องพักรองคณบดีฝ่ายบริหาร"),
    (r"ห้องพักอาจารย์", "ห้องพักอาจารย์"),
    (r"ห้องประชุม\s*(\d*)", "ห้องประชุม"),
    (r"ห้องสมุด", "ห้องสมุดคณะ"),
    (r"ห้องอ่านหนังสือ", "ห้องอ่านหนังสือ"),
    (r"ห้องถ่ายเอกสาร|ร้านถ่ายเอกสาร", "ห้องถ่ายเอกสาร"),
    (r"ห้องคอมพิวเตอร์", "ห้องคอมพิวเตอร์"),
    (r"ห้องเรียน", "ห้องเรียน"),
    (r"ห้องโถง|โถง", "ห้องโถง"),
    (r"ศูนย์พลัง.?งาน", "ศูนย์พลังงาน"),
    (r"ศูนย์ทดสอบประตู", "ศูนย์ทดสอบประตู หน้าต่าง และระบบผนังกระจก"),
    (r"หน่วยบริการสารสนเทศ|หน่วย\s*it", "หน่วยบริการสารสนเทศ (IT)"),
    (r"หน่วยวิเทศ", "หน่วยวิเทศสัมพันธ์"),
    (r"งานคลังและพัสดุ", "งานคลังและพัสดุ"),
    (r"work\s*shop|shop", "Workshop"),
    (r"บันได", "บันได"),
    (r"ทางเชื่อม", "ทางเชื่อมระหว่างอาคาร"),
    (r"ทางเดิน", "ทางเดิน"),
    (r"ระเบียง", "ระเบียง"),
    (r"พื้นอิฐตัวหนอน", "พื้นอิฐตัวหนอน รอบอาคาร"),
    (r"ถนน", "ถนน"),
    (r"หน้าอาคาร", "หน้าอาคาร"),
    (r"หลังอาคาร", "หลังอาคาร"),
    (r"รอบอาคาร", "รอบอาคาร"),
    (r"ทั้งอาคาร", "ทั้งอาคาร"),
    (r"ภาควิชา\s*วิศวกรรม\s*(?:โยธา|เครื่องกล|อุตสาหการ|คอมพิวเตอร์|การอาหาร|เกษตร|ชลประทาน)",
     None),  # ใช้ข้อความภาควิชาตามที่พิมพ์ (จัดรูปในโค้ด)
]

# ชื่อพื้นที่ที่ "มีที่เดียว" ใช้ระบุสถานที่ได้เองแม้ไม่มีเลขอาคาร/ห้อง
# (ชื่อทั่วไป เช่น ห้องน้ำ/ห้อง Server/ห้องธุรการ ต้องมีอาคารหรือรหัสห้องประกอบ ไม่งั้นไม่รู้ว่าที่ไหน)
UNIQUE_AREAS = {
    "ห้องงานบริการการศึกษาและพัฒนานิสิต", "สำนักประสานงานโลจิสติกส์", "สำนักงานเลขานุการ", "ห้องสมุดคณะ",
    "ห้องพักรองคณบดีฝ่ายบริหาร", "ศูนย์พลังงาน", "ศูนย์ทดสอบประตู หน้าต่าง และระบบผนังกระจก",
    "หน่วยบริการสารสนเทศ (IT)", "หน่วยวิเทศสัมพันธ์", "งานคลังและพัสดุ",
}

# อาการเสีย/ไม่ใช่สถานที่ -> ถ้าไม่มีอาคาร/ห้อง/ชื่อพื้นที่ ให้ถือว่า invalid
PROBLEM_WORDS = r"ไม่|เปิด|เสีย|แอร์|เครื่อง|ระบบ|เบอร์|สัญญาณ|รถตู้|ทะเบียน|erp|ล็อก|ขัด|รั่ว|ดับ"

TYPOS = [
    (r"[\u200b\u200c\u200d\ufeff]", ""),
    (r"อาคารื", "อาคาร"),
    (r"ตึก", "อาคาร"),
    (r"(?<![ก-๙])อ\.\s*(?=\d)", "อาคาร "),
    (r"ช้น|ชั้ั", "ชั้น"),
    (r"หัองน้ำ", "ห้องน้ำ"),
    (r"ห้อ\.อ\.[ก-๙]+", ""),
    (r"ศูนย์พลังาน", "ศูนย์พลังงาน"),
    (r"บริการการศึกษษ", "บริการการศึกษา"),
    (r"(?<=[ก-๙])(?=E\s?\d)", " "),       # เรียนE8302 -> เรียน E8302
]


@dataclass
class Place:
    raw: str = ""
    valid: bool = False
    buildings: list = field(default_factory=list)  # ['อาคาร 9', 'อาคารกิจการนิสิต']
    floors: list = field(default_factory=list)     # [2] / [3, 4]
    rooms: list = field(default_factory=list)      # ['9203']
    area: str = ""
    reason: str = ""                               # ทำไม invalid / ตัดอะไรทิ้ง
    dropped: str = ""                              # ข้อความส่วนเกินที่ตัดออกไป

    @property
    def text(self):
        parts = []
        if self.buildings:
            parts.append(", ".join(self.buildings))
        if self.floors:
            parts.append("ชั้น " + ", ".join(map(str, self.floors)))
        if self.rooms:
            parts.append("ห้อง " + ", ".join(self.rooms))
        s = " ".join(parts)
        if self.area:
            s = f"{s} ({self.area})" if self.rooms else (f"{s} {self.area}" if s else self.area)
        return s


def _nums(seq):
    return [int(x) for x in re.findall(r"\d+", seq)]


def parse_place(raw):
    p = Place(raw="" if raw is None else str(raw))
    t = re.sub(r"\s+", " ", p.raw.replace("\xa0", " ")).strip()
    if not t:
        p.reason = "ว่าง"
        return p
    if re.search(r"รถตู้|ทะเบียน", t):
        p.reason = "เป็นยานพาหนะ ไม่ใช่สถานที่"
        return p
    for pat, rep in TYPOS:
        t = re.sub(pat, rep, t, flags=re.I)
    t = re.sub(r"(อาคาร|ชั้น|ห้อง)\s*(?=[0-9A-Za-z])", r"\1 ", t)
    t = re.sub(r"(?<=\d)(?=ชั้น|ห้อง)", " ", t)

    # ---- อาคารที่เรียกด้วยชื่อ
    named = []
    for pat, std in NAMED_BUILDINGS:
        if re.search(pat, t, flags=re.I):
            if std not in named:
                named.append(std)
            t = re.sub(pat, " ", t, flags=re.I)

    # ---- อาคารที่เป็นแค่จุดอ้างอิง (ติด/ข้างอาคาร N) ไม่ใช่ที่เกิดเหตุ -> ตัดออกจากรายการอาคาร
    landmark = re.findall(r"(?:ติด|ข้าง|ใกล้)\s*อาคาร\s*\d+", t)
    t = re.sub(r"(?:ติด|ข้าง|ใกล้)\s*อาคาร\s*\d+", " ", t)

    # ---- อาคารเลข
    bnums = []
    for m in re.finditer(r"อาคาร\s*(\d+)", t):
        n = int(m.group(1))
        if n not in VALID_BUILDINGS:
            p.reason = f"อาคาร {n} ไม่มีจริง"
            return p
        if n not in bnums:
            bnums.append(n)
    t = re.sub(r"อาคาร\s*\d+", " ", t)
    t = re.sub(r"อาคาร(?:ปฏิบัติการ)?", " ", t)

    # ---- ชั้น
    floors = []
    if re.search(r"ชั้นล่าง", t):
        floors.append(1)
        t = t.replace("ชั้นล่าง", " ")
    for m in re.finditer(r"ชั้น\s*(\d+(?:\s*(?:,|และ|-|/)\s*\d+)*)", t):
        vals = _nums(m.group(1))
        if any(v > 9 for v in vals):          # เช่น "ชั้น 9501" = ใส่เลขห้องผิดช่อง -> ปล่อยให้ส่วนรหัสห้องจัดการ
            t = t.replace(m.group(0), " " + " ".join(map(str, vals)) + " ", 1)
            continue
        for v in vals:
            if v not in floors:
                floors.append(v)
        t = t.replace(m.group(0), " ", 1)
    t = re.sub(r"ชั้น", " ", t)

    # ---- รหัสห้อง 4 หลัก (E9203 -> อาคาร 9 ชั้น 2 ห้อง 03)
    rooms, rb, rf = [], [], []
    spans = []
    for m in re.finditer(r"(?<![\dA-Za-zก-๙])E?\s?(\d{4})(?!\d)", t):
        c = m.group(1)
        b, f = int(c[0]), int(c[1])
        if b not in VALID_BUILDINGS or f == 0:
            continue
        spans.append(m.span())
        if c not in rooms:
            rooms.append(c)
        if b not in rb:
            rb.append(b)
        if f not in rf:
            rf.append(f)
    for i, j in reversed(spans):
        t = t[:i] + " " + t[j:]

    # ---- ห้อง 3 หลัก / 2/4 ในอาคารที่เรียกด้วยชื่อ (ห้อง 204 = ชั้น 2 ห้อง 04)
    if named and not bnums and not rooms:
        m = re.search(r"ห้อง\s*(\d)\s*/\s*(\d{1,2})|ห้อง\s*(\d)(\d{2})(?!\d)", t)
        if m:
            f = int(m.group(1) or m.group(3))
            r = (m.group(2) or m.group(4)).zfill(2)
            rooms.append(f"{f}{r}")
            if f not in floors:
                floors.append(f)
            t = t.replace(m.group(0), " ", 1)

    # ---- ตรวจความขัดแย้งกับรหัสห้อง
    if rooms:
        if bnums and any(b not in bnums for b in rb):
            p.reason = f"อาคาร {bnums} ขัดกับรหัสห้อง {rooms}"
            return p
        if floors and any(f not in floors for f in rf) and not named:
            p.reason = f"ชั้น {floors} ขัดกับรหัสห้อง {rooms}"
            return p
        for b in rb:
            if b not in bnums:
                bnums.append(b)
        for f in rf:
            if f not in floors:
                floors.append(f)

    p.buildings = [f"อาคาร {b}" for b in sorted(bnums)] + named
    p.floors = sorted(floors)
    p.rooms = rooms

    # ---- ชื่อพื้นที่ (ส่วนที่เหลือ)
    rest = re.sub(r"\([^)]*\)", lambda m: " " + m.group(0)[1:-1] + " ", t)
    rest = re.sub(r"[\s,/&]+", " ", rest).strip()
    area = ""
    for pat, std in AREAS:
        m = re.search(pat, rest, flags=re.I)
        if not m:
            continue
        if std is None:                                   # ชื่อภาควิชา
            area = re.sub(r"\s+", "", m.group(0)).replace("ภาควิชา", "ภาควิชา")
        elif std == "ห้องประชุม" and m.group(1):
            area = f"ห้องประชุม {m.group(1)}"
        else:
            area = std
        rest = (rest[:m.start()] + " " + rest[m.end():]).strip()
        break
    leftover = re.sub(r"\b(?:ห้อง|ที่|ของ|และ|ใน|ด้าน|หน้า|ฝั่ง|ฯ|\d+)\b|[ฯ]", " ", rest)
    leftover = re.sub(r"\s+", " ", leftover).strip(" ,.-")
    p.area = area

    has_loc = bool(p.buildings or p.rooms)
    if not has_loc and not area:
        if re.search(PROBLEM_WORDS, p.raw, flags=re.I):
            p.reason = "เป็นอาการ/ปัญหา ไม่ใช่สถานที่"
        else:
            p.reason = "ไม่พบอาคาร/ห้อง/ชื่อพื้นที่ที่ระบุได้"
        return p

    if not has_loc and area not in UNIQUE_AREAS:
        p.reason = f"ระบุอาคารไม่ได้ ('{area}' มีหลายที่ ต้องมีอาคารหรือรหัสห้อง)"
        return p

    p.valid = True
    if landmark:
        p.dropped = " ".join(landmark + ([p.dropped] if p.dropped else []))
    if leftover and re.search(PROBLEM_WORDS, leftover, flags=re.I):
        p.dropped = leftover       # อาการ/คำอธิบายที่ปนมาในช่องสถานที่ ไม่เก็บใน place
    return p


log = []  # (sheet, row, field, original, cleaned, rule)


def L(sheet, row, field, orig, new, rule):
    log.append((sheet, row, field, "" if orig is None else str(orig), "" if new is None else str(new), rule))


def blank(v):
    return v is None or (isinstance(v, float) and pd.isna(v)) or (isinstance(v, str) and not v.strip())


def txt(v):
    """str เดียว: NaN->None, strip, ยุบ whitespace/newline เป็นช่องว่างเดียว"""
    if blank(v):
        return None
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


# ---------- extract ----------
def extract(src):
    x = pd.ExcelFile(src)
    frames = []
    # sheet 1: header อยู่แถว 1 (คอลัมน์ A..K) แต่ข้อมูลเริ่มแถว 4 และเลื่อนไปขวา 1 คอลัมน์ (B..L)
    d = x.parse("การตอบแบบฟอร์ม 1", header=None, skiprows=3).iloc[:, 1:12]
    d.columns = ["ts", "name", "status", "phone", "type", "place", "detail", "reported", "extra", "dept", "img"]
    d["sheet"], d["row"] = "ฟอร์ม 1", range(4, 4 + len(d))
    frames.append(d)
    # sheet 3, 4: header แถว 1 ลำดับคอลัมน์เดียวกัน
    for name, tag in [("การตอบแบบฟอร์ม 3", "ฟอร์ม 3"), ("การตอบแบบฟอร์ม 4", "ฟอร์ม 4")]:
        d = x.parse(name)
        d.columns = RAW
        d["sheet"], d["row"] = tag, range(2, 2 + len(d))
        frames.append(d)
    return pd.concat(frames, ignore_index=True)


# ---------- transform helpers ----------
def parse_reported(v):
    """คืน datetime หรือ None; แปลงปี พ.ศ./ปี 2 หลัก/สตริง m/d/yyyy"""
    if blank(v):
        return None
    if isinstance(v, str):
        m = re.match(r"^\s*(\d{1,2})/(\d{1,2})/(\d{1,4})\s+(\d{1,2}):(\d{2})", v)
        if not m:
            return None
        mo, dd, yy, hh, mi = map(int, m.groups())
        yy = yy + 2500 if yy < 100 else yy
        v = datetime(yy, mo, dd, hh, mi)
    v = pd.Timestamp(v).to_pydatetime()
    return v


def be_to_ce(dt):
    return dt.replace(year=dt.year - 543) if dt.year > 2400 else dt


def norm_cmp(s):
    """รูปแบบสำหรับเทียบความซ้ำ: ตัดเว้นวรรค/เครื่องหมาย/ๆ/ค่ะ/ครับ ท้ายประโยค"""
    s = re.sub(r"[\s\W_]+", "", (s or "").lower())
    return re.sub(r"(ค่ะ|คะ|ครับ|นะคะ|นะครับ)+$", "", s)


def clean_type(t, sheet, row):
    t0 = txt(t)
    if t0 in STD_TYPES:
        return t0
    if t0 in TYPE_MAP:
        return TYPE_MAP[t0]
    return "งานอื่นๆ"


def clean_status(s):
    s0 = txt(s)
    if s0 in STATUS_STD:
        return s0
    return STATUS_MAP.get(s0, "อื่นๆ")


def is_test(r):
    name = txt(r["name"]) or ""
    if name.startswith("ทดสอบ"):
        return True
    for c in ("type", "place", "detail", "extra"):
        if "ทดสอบ" in (txt(r[c]) or "").replace("ศูนย์ทดสอบ", ""):  # "ศูนย์ทดสอบ..." เป็นชื่อหน่วยงาน
            return True
    return False


# ---------- SQL / dim_date helpers ----------
FISCAL_YEAR_BE = True   # dim_date.fiscal_year: True = ปีงบประมาณ พ.ศ. (เริ่ม 1 ต.ค.), False = ค.ศ.
THAI_DAYS = ["จันทร์", "อังคาร", "พุธ", "พฤหัสบดี", "ศุกร์", "เสาร์", "อาทิตย์"]
THAI_MONTHS = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
               "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]


def q(v):
    """ค่า -> SQL literal (MySQL)"""
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return "NULL"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return str(int(v)) if float(v).is_integer() else str(v)
    return "'" + str(v).replace("\\", "\\\\").replace("'", "''") + "'"


def date_rows(d0, d1):
    """ทุกวันของปีปฏิทินที่ครอบช่วงข้อมูล"""
    d = datetime(d0.year, 1, 1)
    end = datetime(d1.year, 12, 31)
    while d <= end:
        fy = d.year + (1 if d.month >= 10 else 0)
        if FISCAL_YEAR_BE:
            fy += 543
        yield (int(d.strftime("%Y%m%d")), d.strftime("%Y-%m-%d"), THAI_DAYS[d.weekday()], d.month,
               THAI_MONTHS[d.month - 1], (d.month - 1) // 3 + 1, d.year, fy)
        d += timedelta(days=1)


def run_excel(src, out_prefix):
    """Excel (Google Form ใบแจ้งซ่อมเดิม) -> DataFrame 9 คอลัมน์ + repair_requests_clean.xlsx/.csv"""
    raw = extract(src)
    out = []
    dropped = []
    rejected = []
    for _, r in raw.iterrows():
        sh, row = r["sheet"], r["row"]
        if is_test(r):
            L(sh, row, "(ทั้งแถว)", f"{txt(r['name'])} | {txt(r['type'])} | {txt(r['place'])} | {txt(r['detail'])}", "", "ลบ: แถวทดสอบระบบ")
            dropped.append("test")
            continue

        rec = {"_sheet": sh, "_row": row, "_name": txt(r["name"]), "_ts": pd.Timestamp(r["ts"]).to_pydatetime()}

        # สังกัด
        dept = txt(r["dept"]) or "ไม่ระบุ"
        rec["สังกัด"] = dept

        # สถานภาพ
        st0 = txt(r["status"])
        st = clean_status(st0)
        if st != st0:
            L(sh, row, "สถานภาพผู้แจ้งซ่อม", st0, st, "จัดเข้าชุดสถานภาพมาตรฐาน (เจ้าหน้าที่/อาจารย์/นิสิต/ช่างเทคนิค/แม่บ้าน/อื่นๆ)")
        rec["สถานภาพผู้แจ้งซ่อม"] = st

        # ประเภท
        t0 = txt(r["type"])
        t = clean_type(t0, sh, row)
        if t != t0:
            L(sh, row, "ประเภทงานซ่อมบำรุง", t0, t, "จัดข้อความที่พิมพ์เองเข้าประเภทมาตรฐาน")
        rec["ประเภทงานซ่อมบำรุง"] = t

        # สถานที่: แยกเป็น อาคาร/ชั้น/ห้อง/ชื่อพื้นที่ (E9203 = อาคาร 9 ชั้น 2 ห้อง 9203)
        pl = parse_place(r["place"])
        if not pl.valid:
            rejected.append({"แหล่งที่มา": sh, "แถวในไฟล์เดิม": row, "แจ้งซ่อมเมื่อ(ประทับเวลา)": rec["_ts"],
                             "ประเภทงาน": txt(r["type"]), "สถานที่ที่กรอก": txt(r["place"]),
                             "รายละเอียด": txt(r["detail"]), "เหตุผลที่ใช้ไม่ได้": pl.reason})
            L(sh, row, "สถานที่เกิดปัญหา", txt(r["place"]), "", f"สถานที่ใช้ไม่ได้: {pl.reason}" + (" -> ลบแถว" if DROP_INVALID_PLACE else " -> เว้นว่าง"))
            if DROP_INVALID_PLACE:
                dropped.append("place")
                continue
        if pl.valid:
            if pl.text != txt(r["place"]):
                L(sh, row, "สถานที่เกิดปัญหา", txt(r["place"]), pl.text, "แยก/จัดรูป อาคาร-ชั้น-ห้อง (รหัสห้อง 4 หลัก: อาคาร·ชั้น·เลขห้อง, ตัด E)")
            if pl.dropped:
                L(sh, row, "สถานที่เกิดปัญหา", pl.dropped, "", "ตัดส่วนที่เป็นอาการ/คำอธิบายออกจากช่องสถานที่")
        rec["สถานที่เกิดปัญหา"] = pl.text if pl.valid else None
        rec.update(place_columns(pl))
        d0 = r["detail"]
        d = txt(d0)
        if d != (None if blank(d0) else str(d0)):
            L(sh, row, "รายละเอียดของปัญหา", d0, d, "trim/ยุบช่องว่าง/ขึ้นบรรทัดใหม่เป็นช่องว่าง")

        # รวม รายละเอียดเพิ่มเติม เข้า รายละเอียดของปัญหา (ตัดค่าว่าง/ซ้ำ/ที่ถูกรวมอยู่แล้ว)
        merged, rule = merge_detail(d, r["extra"])
        if rule:
            L(sh, row, "รายละเอียดของปัญหา (รวม)", f"{d} || {txt(r['extra'])}", merged, rule)
        rec["รายละเอียดของปัญหา"] = merged or "ไม่ระบุ"

        # แจ้งซ่อมเมื่อ
        ts_min = rec["_ts"].replace(second=0, microsecond=0)
        rep0 = r["reported"]
        rep = parse_reported(rep0)
        rule = None
        if rep is not None:
            fixed = be_to_ce(rep)
            if fixed != rep:
                rule = "แปลงปี พ.ศ. -> ค.ศ."
            rep = fixed
        if rep is None:
            rep, rule = ts_min, "อ่านวันที่ไม่ได้ -> ใช้ประทับเวลา"
        elif rep > ts_min + DATE_FUTURE_TOLERANCE:
            rep, rule = ts_min, "วันที่อยู่หลังเวลาส่งฟอร์ม (พิมพ์ผิด) -> ใช้ประทับเวลา"
        elif ts_min - rep > DATE_FALLBACK_AFTER:
            rep, rule = ts_min, "วันที่เก่ากว่าเวลาส่งฟอร์มเกิน 7 วัน (พิมพ์ผิด ปี/เดือน/วัน) -> ใช้ประทับเวลา"
        if rule:
            L(sh, row, "แจ้งซ่อมเมื่อ", rep0, rep.strftime("%Y-%m-%d %H:%M"), rule)
        rec["แจ้งซ่อมเมื่อ"] = rep
        out.append(rec)

    df = pd.DataFrame(out)

    # ลบซ้ำ: ผู้แจ้งคนเดียวกัน เนื้อหาเหมือนกันทุกช่อง ส่งซ้ำภายใน 24 ชม.
    df = df.sort_values(["_ts"]).reset_index(drop=True)
    keep = []
    last = {}
    for i, r in df.iterrows():
        key = (r["_name"], r["สังกัด"], r["ประเภทงานซ่อมบำรุง"], r["สถานที่เกิดปัญหา"], r["รายละเอียดของปัญหา"])
        if key in last and r["_ts"] - last[key][0] <= timedelta(hours=24):
            L(r["_sheet"], r["_row"], "(ทั้งแถว)", f"{r['สถานที่เกิดปัญหา']} | {r['รายละเอียดของปัญหา']}",
              "", f"ลบ: ส่งซ้ำภายใน 24 ชม. (ซ้ำกับ {last[key][1]} แถว {last[key][2]})")
            dropped.append("dup")
            continue
        last[key] = (r["_ts"], r["_sheet"], r["_row"])
        keep.append(i)
    df = df.loc[keep]

    df = df.sort_values("แจ้งซ่อมเมื่อ").reset_index(drop=True)
    final = df[TARGET].copy()

    # ---------- load (ไฟล์ clean ของ Excel) ----------
    final.assign(**{"แจ้งซ่อมเมื่อ": final["แจ้งซ่อมเมื่อ"].dt.strftime("%Y-%m-%d %H:%M")}) \
        .to_csv(out_prefix + ".csv", index=False, encoding="utf-8-sig")

    wb = Workbook()
    ws = wb.active
    ws.title = "repair_clean"
    ws.append(TARGET)
    for rec in final.itertuples(index=False):
        ws.append([v if not (isinstance(v, float) and pd.isna(v)) else None for v in rec])
    style_sheet(ws, [18, 28, 16, 28, 16, 8, 16, 32, 70], wrap_cols=(8, 9), date_cols=(1,), autofilter=True)

    lg = wb.create_sheet("cleaning_log")
    lg.append(["แหล่งที่มา", "แถวในไฟล์เดิม", "ฟิลด์", "ค่าเดิม", "ค่าใหม่", "กฎที่ใช้"])
    for e in sorted(log, key=lambda e: (e[0], e[1])):
        lg.append(list(e))
    style_sheet(lg, [12, 14, 24, 50, 40, 70], wrap_cols=range(1, 7), autofilter=True)

    rj = wb.create_sheet("rejected_place")
    rcols = ["แหล่งที่มา", "แถวในไฟล์เดิม", "แจ้งซ่อมเมื่อ(ประทับเวลา)", "ประเภทงาน", "สถานที่ที่กรอก", "รายละเอียด", "เหตุผลที่ใช้ไม่ได้"]
    rj.append(rcols)
    for rr in rejected:
        rj.append([rr[c] for c in rcols])
    style_sheet(rj, [12, 14, 22, 28, 40, 50, 50], wrap_cols=range(1, 8), date_cols=(3,))

    wb.save(out_prefix + ".xlsx")
    print(f"[excel] raw {len(raw)} | ลบ: ทดสอบ {dropped.count('test')}, ซ้ำ {dropped.count('dup')}, "
          f"สถานที่ใช้ไม่ได้ {dropped.count('place')} | เหลือ {len(final)}")
    return final


def place_columns(pl):
    """Place -> คอลัมน์ อาคาร/ชั้น/ห้อง/ชื่อพื้นที่ (หลายค่าคั่นด้วย ', ')"""
    return {"อาคาร": ", ".join(pl.buildings) or None,
            "ชั้น": ", ".join(map(str, pl.floors)) or None,
            "ห้อง": ", ".join(pl.rooms) or None,
            "ชื่อพื้นที่": pl.area or None}


def merge_detail(d, extra):
    """รวม รายละเอียดเพิ่มเติม เข้า รายละเอียดของปัญหา -> (ข้อความ, กฎที่ใช้ หรือ None)"""
    e = txt(extra)
    if e is not None and e.lower() in NULL_EXTRA:
        return d, "รายละเอียดเพิ่มเติมเป็นค่าว่างความหมาย (-, ไม่มี ฯลฯ) -> ตัดทิ้ง"
    if e is None:
        return d, None
    nd, ne = norm_cmp(d), norm_cmp(e)
    if not d:
        return e, "รายละเอียดของปัญหาว่าง -> ใช้รายละเอียดเพิ่มเติมแทน"
    if ne == nd or ne in nd:
        return d, "รายละเอียดเพิ่มเติมซ้ำ/อยู่ในรายละเอียดของปัญหาแล้ว -> เก็บอันเดียว"
    if nd in ne:
        return e, "รายละเอียดของปัญหาอยู่ในรายละเอียดเพิ่มเติมแล้ว -> เก็บอันที่ยาวกว่า"
    return f"{d} | {e}", "รวมรายละเอียดของปัญหา + รายละเอียดเพิ่มเติม (คั่นด้วย |)"


def style_sheet(ws, widths, wrap_cols=(), date_cols=(), autofilter=False):
    font = Font(name="Arial", size=10)
    hfont = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="1F4E78")
    wrap_cols = set(wrap_cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.font = font
            c.alignment = Alignment(vertical="top", wrap_text=c.column in wrap_cols)
    for c in ws[1]:
        c.font, c.fill = hfont, hfill
    for col in date_cols:
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col).number_format = "yyyy-mm-dd hh:mm"
    ws.freeze_panes = "A2"
    if autofilter:
        ws.auto_filter.ref = ws.dimensions


# =====================================================================
# แหล่งที่ 2: API ระบบครุภัณฑ์ (CPEAssets)
# =====================================================================
API_BASE = "https://cpe.eng.kps.ku.ac.th/CPEAssets/assets_api/"
# endpoint -> "api" (เรียก API จริง) | "mock" (อ่าน mock_api/<endpoint>.json)
#           | "file:<path>" (ไฟล์ข้อมูลจริงที่ export ไว้ใน repo, path นับจากโฟลเดอร์สคริปต์; JSON เป็น list หรือ {data: [...]})
# เมื่อระบบจริงเปิด endpoint ไหนแล้ว ให้เปลี่ยนเป็น "api" ได้เลย รูปแบบ JSON เหมือนกัน
HERE = os.path.dirname(os.path.abspath(__file__))
SOURCES = {
    "assets": "api",
    "asset_types": "mock",
    "asset_categories": "mock",
    "departments": "mock",
    "users": "mock",
    "locations": "mock",
    "maintenance_type": "mock",
    "maintenance_asset": "mock",
    "maintenance_general": "mock",
    "borrow": "mock",
    "borrow_request_items": "mock",
}
# endpoint ที่เป็น "api": ถ้าเรียกไม่ได้ (เน็ตหลุด / server ล่ม / ตอบ success=false) ใช้ไฟล์ export นี้แทน
FALLBACK_FILES = {"assets": "data/assets_real.json"}
USED = {}   # endpoint -> แหล่งที่ใช้จริงรอบนี้: "api" | "cache" (--offline) | "file:<path>" | "mock"

# ชื่อสังกัดในฟอร์ม Excel -> ชื่อใน departments ของระบบครุภัณฑ์ (ที่เหลือชื่อตรงกันอยู่แล้ว)
EXCEL_DEPT_ALIAS = {
    "สำนักงานเลขานุการ": "สำนักงานเลขานุการคณะ",
    "ภาควิชาวิศวกรรมการอาหาร": "ภาควิชาวิศวกรรมอาหาร",
}

# แหล่งเงิน: 8 แบบที่สะกดใน API -> 4 กลุ่ม
SOURCE_GROUP = {
    "เงินรายได้": "เงินรายได้", "รายได้": "เงินรายได้", "เงินรายได้ภาคฯ": "เงินรายได้",
    "เงินรายได้โครงการฯ": "เงินรายได้", "งบประมาณ": "เงินงบประมาณ", "เงินงบประมาณ": "เงินงบประมาณ",
    "โครงการภาคพิเศษ": "โครงการภาคพิเศษ", "บริจาค": "บริจาค",
}

def read_file(path):
    """ไฟล์ข้อมูลใน repo (path นับจากโฟลเดอร์สคริปต์): JSON เป็น list หรือ {data: [...]}"""
    with open(os.path.join(HERE, path), encoding="utf-8") as f:
        body = json.load(f)
    return body["data"] if isinstance(body, dict) else body


def fetch_api(endpoint):
    rows, page, total = [], 1, None
    while True:
        url = f"{API_BASE}{endpoint}.php?limit=100&page={page}"
        req = urllib.request.Request(url, headers={"User-Agent": "etl-warehouse"})
        with urllib.request.urlopen(req, timeout=60) as resp:
            body = json.loads(resp.read())
        if not body.get("success", True):
            raise RuntimeError(f"{url}: {body.get('message')}")
        total = body.get("total", total)
        rows += body["data"]
        if not body["data"] or (total is not None and len(rows) >= total):
            return rows
        page += 1


def fetch(endpoint, mock_dir, raw_dir, offline=False):
    """ดึงทุกหน้าของ endpoint -> list ของ record; เก็บ raw ไว้ที่ raw_dir/<endpoint>.json
    endpoint ที่มีใน FALLBACK_FILES: ถ้าเรียก API (หรือ raw ตอน --offline) ไม่ได้ ใช้ไฟล์นั้นแทน"""
    src = SOURCES[endpoint]
    if src == "mock":
        USED[endpoint] = "mock"
        with open(os.path.join(mock_dir, f"{endpoint}.json"), encoding="utf-8") as f:
            return json.load(f)["data"]
    if src.startswith("file:"):
        USED[endpoint] = src
        return read_file(src[len("file:"):])
    cache = os.path.join(raw_dir, f"{endpoint}.json")
    try:
        if offline:
            with open(cache, encoding="utf-8") as f:
                USED[endpoint] = "cache"
                return json.load(f)["data"]
        rows = fetch_api(endpoint)
    except (OSError, ValueError, RuntimeError) as e:      # URLError/timeout เป็น OSError, JSON เสียเป็น ValueError
        fallback = FALLBACK_FILES.get(endpoint)
        if not fallback:
            raise
        print(f"[warn] {endpoint}: {'อ่าน raw' if offline else 'เรียก API'} ไม่ได้ ({e}) -> ใช้ {fallback} แทน")
        USED[endpoint] = "file:" + fallback
        return read_file(fallback)
    with open(cache, "w", encoding="utf-8") as f:
        json.dump({"success": True, "total": len(rows), "fetched_at": datetime.now().isoformat(timespec="seconds"),
                   "data": rows}, f, ensure_ascii=False)
    USED[endpoint] = "api"
    return rows


def dt(v):
    return None if v in (None, "") else datetime.strptime(str(v)[:19], "%Y-%m-%d %H:%M:%S" if len(str(v)) > 10 else "%Y-%m-%d")


def dkey(v):
    return None if v is None else int(v.strftime("%Y%m%d"))


# ปีงบที่ได้มา ถอดจากท้ายรหัสครุภัณฑ์: ".../001-68" หรือ ".../ร.54-1" -> ปีงบ 2568 / 2554
ASSET_CODE_FY = [re.compile(r"/\d{3}-(\d{2})\s*$"), re.compile(r"/[ก-๙]{1,3}\.?-?(\d{2})-\d+\s*$")]
BELOW_THRESHOLD_PREFIX = ("FN", "FP", "FU")   # ต่ำกว่าเกณฑ์ / ต่ำกว่าเกณฑ์จากบริจาค / ต่ำกว่าเกณฑ์รับโอน


def fy_from_code(code):
    for pat in ASSET_CODE_FY:
        m = pat.search(code or "")
        if m and 2530 <= 2500 + int(m.group(1)) <= 2600:
            return 2500 + int(m.group(1))
    return None


def category_code_of(type_code):
    """FA05 / FN05 / FT05 -> "05" (เลขหมวดใน type_code)"""
    m = re.fullmatch(r"[A-Z]{2}(\d{2})", (type_code or "").strip())
    return m.group(1) if m else None


def acquisition_date(a):
    """วันที่ได้มา: purchase_date > ปีงบจากรหัส (1 ต.ค. ต้นปีงบ) > year_purchase (1 ม.ค.)"""
    if a.get("purchase_date"):
        return dt(a["purchase_date"])
    fy = fy_from_code(a.get("asset_code"))
    if fy:
        return datetime(fy - 544, 10, 1)
    y = txt(a.get("year_purchase"))
    return datetime(int(y), 1, 1) if y and y.isdigit() and 1950 <= int(y) <= 2100 else None


def load_history(history_dir):
    """snapshot ครุภัณฑ์ที่เคยดึงไว้ทุกครั้ง: {datetime: [asset, ...]} จากไฟล์ assets_YYYYMMDD.json"""
    snaps = {}
    for f in sorted(os.listdir(history_dir)):
        m = re.fullmatch(r"assets_(\d{8})\.json", f)
        if m:
            with open(os.path.join(history_dir, f), encoding="utf-8") as fh:
                snaps[datetime.strptime(m.group(1), "%Y%m%d")] = json.load(fh)["data"]
    return snaps



class Dim:
    """dimension ที่ให้ surrogate key ตามลำดับที่เพิ่ม; lookup ด้วย natural key"""

    def __init__(self, cols):
        self.cols, self.rows, self.index = cols, [], {}

    def add(self, nk, **vals):
        if nk not in self.index:
            self.index[nk] = len(self.rows) + 1
            self.rows.append({self.cols[0]: self.index[nk], **vals})
        return self.index[nk]

    def get(self, nk):
        return self.index.get(nk)


def build_warehouse(excel, api, snapshots, as_of):
    """รวม 2 แหล่ง -> dict ของตาราง warehouse (list of dict) + รายการปัญหาที่พบ"""
    notes = []
    # ---------- dim_department ----------
    # แถวที่ถูกลบ (is_deleted = 1) เก็บไว้เฉพาะที่ยังมีข้อมูลอ้างถึง เพื่อไม่ให้ประวัติเดิมหาย
    dep_used = {r.get("department_id") for ep in ("assets", "users", "borrow", "maintenance_general") for r in api[ep]}
    dep_used |= {a.get("department_id") for rows in snapshots.values() for a in rows}
    dep = Dim(["department_key", "department_id", "department_name"])
    dep_by_id, name_to_key = {}, {}
    for d in api["departments"]:
        if d.get("is_deleted") and d["department_id"] not in dep_used:
            continue
        k = dep.add(("id", d["department_id"]), department_id=d["department_id"], department_name=txt(d["department_name"]))
        dep_by_id[d["department_id"]] = k
        for n in (txt(d["department_name"]), txt(d.get("shortname"))):
            if n:
                name_to_key.setdefault(n, k)

    def dep_from_excel(name):
        """Excel ไม่มี department_id -> จับคู่ด้วยชื่อ (alias / ชื่อเต็ม / shortname); ไม่เจอ = แถวใหม่ department_id NULL"""
        n = EXCEL_DEPT_ALIAS.get(name, name)
        return name_to_key.get(n) or dep.add(("name", n), department_id=None, department_name=n)

    # ---------- dim_user ----------
    usr = Dim(["user_key", "user_id", "fullname", "role_status"])
    for u in api["users"]:
        usr.add(u["user_id"], user_id=u["user_id"], fullname=txt(u["fullname"]), role_status=txt(u.get("role_status")))

    # ---------- dim_maintenance_type ----------
    mt = Dim(["maintenance_type_key", "type_id", "type_name"])
    mt_by_id = {}
    for t in api["maintenance_type"]:
        mt_by_id[t["type_id"]] = mt.add(txt(t["type_name"]), type_id=t["type_id"], type_name=txt(t["type_name"]))

    # ---------- dim_location (ใช้ร่วม: ห้องของครุภัณฑ์ + สถานที่แจ้งซ่อม) ----------
    loc = Dim(["location_key", "location_id", "building", "floor", "room", "area_name"])

    def loc_key(b, f, r, a, location_id=None):
        """มีรหัสห้อง: 1 ห้อง = 1 แถว (ชื่อพื้นที่/location_id ใช้ค่าแรกที่เจอ ซึ่งมาจากตาราง locations ก่อน)
        location_id = NULL สำหรับสถานที่ที่มาจากใบแจ้งซ่อมเท่านั้น (ไม่มีในตาราง locations)"""
        if not any([b, f, r, a]):
            return None
        nk = (b, f, r) if r else (b, f, None, a)
        k = loc.add(nk, location_id=location_id, building=b, floor=f, room=r or None, area_name=a)
        row = loc.rows[k - 1]
        if a and not row["area_name"]:
            row["area_name"] = a
        if location_id is not None and row["location_id"] is None:
            row["location_id"] = location_id
        return k

    # ตาราง location ของระบบครุภัณฑ์ (location_name, building, floor, room_detail) กรอกอิสระ
    # -> ผ่านตัวแยกสถานที่ตัวเดียวกับ Excel เพื่อให้ห้องเดียวกันได้ location_key เดียวกัน
    loc_by_id, loc_issue = {}, []
    loc_used = {a["location_id"] for rows in [api["assets"], *snapshots.values()] for a in rows}
    for l in api["locations"]:
        if l.get("is_deleted") and l["location_id"] not in loc_used:
            continue
        b, f, rd, name = txt(l["building"]), txt(l["floor"]), txt(l["room_detail"]), txt(l["location_name"])
        pl = parse_place(" ".join(x for x in [b, f and f"ชั้น {f}", rd and f"ห้อง {rd}", name] if x))
        if pl.valid:
            pc = place_columns(pl)
            area = pc["ชื่อพื้นที่"] or (None if pc["ห้อง"] else name)    # ห้องมีรหัสแล้ว ไม่ต้องใช้ชื่อห้องแยก
            loc_by_id[l["location_id"]] = loc_key(pc["อาคาร"], pc["ชั้น"], pc["ห้อง"], area, l["location_id"])
        else:
            loc_issue.append(f"{l['location_id']} ({pl.reason})")
            loc_by_id[l["location_id"]] = loc_key(b, f, rd, name, l["location_id"])
    if loc_issue:
        notes.append("locations: แยกอาคาร/ชั้น/ห้องไม่ได้ ใช้ค่าตามที่กรอก: " + ", ".join(loc_issue))

    # ---------- dim_asset_category / dim_asset_type / dim_asset ----------
    # asset ใช้รายการล่าสุด + ครุภัณฑ์ที่มีเฉพาะใน snapshot เก่า (ถูกลบไปแล้ว) เพื่อให้ snapshot เก่ายังหา asset_key ได้
    latest = {}
    for d in sorted(snapshots):
        for a in snapshots[d]:
            latest[a["asset_id"]] = a
    for a in api["assets"]:
        latest[a["asset_id"]] = a
    all_assets = [latest[k] for k in sorted(latest)]

    cat = Dim(["category_key", "category_id", "category_name"])
    cat_by_id, cat_by_code = {}, {}
    for c in api["asset_categories"]:
        k = cat.add(c["category_id"], category_id=c["category_id"], category_name=txt(c["category_name_th"]))
        cat_by_id[c["category_id"]] = k
        if c.get("category_code"):
            cat_by_code[txt(c["category_code"])] = k

    typ = Dim(["type_key", "type_id", "type_code", "type_name"])
    type_used = {a["type_id"] for a in all_assets}
    type_code_of = {}
    for t in api["asset_types"]:
        if t.get("is_deleted") and t["type_id"] not in type_used:
            continue
        typ.add(t["type_id"], type_id=t["type_id"], type_code=txt(t["type_code"]), type_name=txt(t["type_name"]))
        type_code_of[t["type_id"]] = txt(t["type_code"])
    for a in all_assets:                                   # ประเภทที่ asset อ้างแต่ asset_types ไม่มี
        if a["type_id"] is not None and typ.get(a["type_id"]) is None:
            typ.add(a["type_id"], type_id=a["type_id"], type_code=txt(a["type_code"]), type_name=txt(a["type_name"]))
            type_code_of[a["type_id"]] = txt(a["type_code"])
    code_to_type = {}
    for tid, code in type_code_of.items():
        code_to_type.setdefault(code, []).append(tid)

    def type_id_of(a):
        """type_id ของ asset; ถ้าว่างแต่มี type_code ที่ตรงกับประเภทเดียว ใช้ประเภทนั้น"""
        if a["type_id"] is not None:
            return a["type_id"]
        ids = code_to_type.get(txt(a.get("type_code")), [])
        return ids[0] if len(ids) == 1 else None

    def category_key_of(a):
        """category_id ของ asset ถ้ามี; ถ้าไม่มี ใช้เลขหมวดใน type_code (FA05 -> 05) จับกับ asset_categories.category_code"""
        if a.get("category_id") in cat_by_id:
            return cat_by_id[a["category_id"]]
        tid = type_id_of(a)
        return cat_by_code.get(category_code_of(txt(a.get("type_code")) or type_code_of.get(tid)))

    ast = Dim(["asset_key", "asset_id", "asset_code", "asset_name", "brand", "model", "year_purchase", "year_fiscal",
               "source", "asset_status", "category_key", "type_key", "location_key"])
    asset_dep = {}
    for a in all_assets:
        src = txt(a["source"])
        fy = txt(a["year_fiscal"]) or (str(fy_from_code(a["asset_code"])) if fy_from_code(a["asset_code"]) else None)
        ast.add(a["asset_id"], asset_id=a["asset_id"], asset_code=txt(a["asset_code"]), asset_name=txt(a["asset_name"]),
                brand=txt(a["brand"]), model=txt(a.get("model")), year_purchase=txt(a["year_purchase"]), year_fiscal=fy,
                source=SOURCE_GROUP.get(src, src), asset_status=txt(a["asset_status"]),
                category_key=category_key_of(a), type_key=typ.get(type_id_of(a)), location_key=loc_by_id.get(a["location_id"]))
        asset_dep[a["asset_id"]] = dep_by_id.get(a["department_id"])
    lost = sorted({a["location_id"] for a in all_assets if a["location_id"] and a["location_id"] not in loc_by_id})
    if lost:
        notes.append(f"dim_asset: location_id ที่ไม่มีใน locations (หรือถูกลบ) -> location_key = NULL: {lost}")
    no_cat = sum(1 for r in ast.rows if r["category_key"] is None)
    if no_cat:
        notes.append(f"dim_asset: {no_cat} รายการหาหมวดไม่ได้ (ไม่มี type_code) -> category_key = NULL")
    if any(r["type_name"] and re.fullmatch(r"[A-Z]{2}\d{2}", r["type_name"]) for r in typ.rows):
        notes.append("dim_asset_type: บาง type_name ในต้นทางเป็นรหัส (เช่น FN08, FN13) ไม่ใช่ชื่อประเภท")

    # ---------- fact_asset_inventory (periodic snapshot: 1 แถว = ครุภัณฑ์ 1 ชิ้น ณ วันที่ดึงข้อมูล) ----------
    fai = []
    for snap in sorted(snapshots):
        for a in snapshots[snap]:
            acq = acquisition_date(a)
            tid = type_id_of(a)
            tcode = txt(a.get("type_code")) or type_code_of.get(tid) or ""
            status = txt(a["asset_status"])
            lk = loc_by_id.get(a["location_id"])
            fai.append({"snapshot_date_key": dkey(snap), "acquisition_date_key": dkey(acq),
                        "asset_key": ast.get(a["asset_id"]), "department_key": dep_by_id.get(a["department_id"]),
                        "type_key": typ.get(tid), "category_key": category_key_of(a), "location_key": lk,
                        "asset_count": 1, "is_available_flag": int(status == "พร้อมใช้งาน"),
                        "is_damaged_flag": int(status in ("ชำรุด", "ส่งซ่อม")),
                        "is_below_threshold_flag": int(tcode.startswith(BELOW_THRESHOLD_PREFIX) or "ต่ำกว่าเกณฑ์" in (txt(a["type_name"]) or "")),
                        "has_location_flag": int(lk is not None),
                        "asset_age_years": max(0, int((snap - acq).days // 365.25)) if acq else None})
    no_acq = sum(1 for r in fai if r["acquisition_date_key"] is None)
    if no_acq:
        notes.append(f"fact_asset_inventory: {no_acq} แถวหาวันที่ได้มาไม่ได้ (ไม่มี purchase_date / ปีในรหัส / year_purchase)")

    # ---------- fact_general_maintenance: Excel (ฟอร์มเดิม) ----------
    fgm = []
    for r in excel.itertuples(index=False):
        rep = r[0].to_pydatetime() if hasattr(r[0], "to_pydatetime") else r[0]
        fgm.append({"reported_date_key": dkey(rep), "completed_date_key": None,
                    "department_key": dep_from_excel(r[1]),
                    "maintenance_type_key": mt.get(r[3]) or mt.add(r[3], type_id=None, type_name=r[3]),
                    "location_key": loc_key(txt(r[4]), txt(r[5]), txt(r[6]), txt(r[7])),
                    "staff_user_key": None, "reporter_status": r[2], "job_status": None, "incident_count": 1,
                    "resolution_hours": None, "resolution_days": None, "_src": "excel"})

    # ---------- fact_general_maintenance: API maintenance_general (ระบบใหม่) ----------
    bad = []
    for g in api["maintenance_general"]:
        pl = parse_place(g["problem_location"])
        if not pl.valid:
            L("API maintenance_general", g["id"], "problem_location", g["problem_location"], "",
              f"สถานที่ใช้ไม่ได้: {pl.reason}" + (" -> ลบแถว" if DROP_INVALID_PLACE else " -> เว้นว่าง"))
            if DROP_INVALID_PLACE:
                bad.append(f"id {g['id']} '{g['problem_location']}' ({pl.reason})")
                continue
        pc = place_columns(pl)
        rep, done = dt(g["reported_at"]), dt(g["completed_at"])
        hours = round((done - rep).total_seconds() / 3600, 2) if rep and done else None
        fgm.append({"reported_date_key": dkey(rep), "completed_date_key": dkey(done),
                    "department_key": dep_by_id.get(g["department_id"]),
                    "maintenance_type_key": mt_by_id.get(g["maintenance_type_id"]),
                    "location_key": loc_key(pc["อาคาร"], pc["ชั้น"], pc["ห้อง"], pc["ชื่อพื้นที่"]) if pl.valid else None,
                    "staff_user_key": usr.get(g["staff_id"]), "reporter_status": clean_status(g["reporter_status"]),
                    "job_status": txt(g["status"]), "incident_count": 1, "resolution_hours": hours,
                    "resolution_days": (done.date() - rep.date()).days if rep and done else None, "_src": "api"})
    if bad:
        notes.append(f"maintenance_general (API): ลบ {len(bad)} แถวที่สถานที่ใช้ไม่ได้: " + "; ".join(bad))
    # Excel = ฟอร์มเดิม, API = ระบบใหม่ที่ใช้ต่อจากนั้น -> ช่วงเวลาไม่ควรซ้อนกัน (ถ้าซ้อน อาจมีงานซ้ำ 2 แหล่ง)
    xl_end = max((r["reported_date_key"] for r in fgm if r["_src"] == "excel"), default=None)
    overlap = sum(1 for r in fgm if r["_src"] == "api" and xl_end and r["reported_date_key"] <= xl_end)
    if overlap:
        notes.append(f"maintenance_general (API): {overlap} แถวแจ้งก่อน/วันเดียวกับ Excel ฉบับสุดท้าย ({xl_end}) อาจซ้ำกับ Excel")

    # ---------- fact_asset_maintenance ----------
    fam = []
    for m in api["maintenance_asset"]:
        rep, done = dt(m["repair_date"]), dt(m["complete_date"])
        cost, fcost = m["cost"], m["followup_cost"]
        total = None if cost is None and fcost is None else round((cost or 0) + (fcost or 0), 2)
        days = m["repair_total_day"] if m["repair_total_day"] is not None else \
            ((done.date() - rep.date()).days if rep and done else None)
        fam.append({"repair_date_key": dkey(rep), "complete_date_key": dkey(done), "asset_key": ast.get(m["asset_id"]),
                    "department_key": asset_dep.get(m["asset_id"]), "approved_by_user_key": usr.get(m["approved_by"]),
                    "vendor_name": txt(m["vendor"]), "repair_status": txt(m["status"]), "maintenance_count": 1,
                    "repair_cost": cost, "followup_cost": fcost, "total_repair_cost": total, "repair_total_days": days,
                    "is_followup_flag": int(bool(m["is_followup"])), "followup_no": m["followup_no"]})

    # ---------- fact_borrow_transaction (grain = ครุภัณฑ์ 1 ชิ้นในใบยืม) ----------
    borrows = {b["id"]: b for b in api["borrow"]}
    fbt = []
    for it in api["borrow_request_items"]:
        b = borrows[it["request_id"]]
        bd, due, act = dt(b["borrow_date"]), dt(b["return_date"]), dt(b["actual_return_date"])
        end = act or (as_of if b["status"] not in ("ปฏิเสธ",) else None)
        overdue = max(0, (end.date() - due.date()).days) if end and due else None
        flag = int(overdue > 0) if overdue is not None else b["is_overdue"]
        fbt.append({"borrow_date_key": dkey(bd), "return_date_key": dkey(due), "actual_return_date_key": dkey(act),
                    "asset_key": ast.get(it["replace_asset_id"] or it["asset_id"]),
                    "borrower_user_key": usr.get(b["user_id"]),
                    "approver_user_key": usr.get(b["approved_by"] or b["teacher_id"]),
                    "department_key": dep_by_id.get(b["department_id"]), "borrow_request_id": b["id"],
                    "item_status": txt(it["item_status"]), "return_condition": txt(b["return_condition"]),
                    "borrow_item_count": 1, "borrow_duration_days": (act.date() - bd.date()).days if act else None,
                    "overdue_days": overdue, "is_overdue_flag": flag,
                    "is_replaced_flag": int(it["replace_asset_id"] is not None)})

    # ---------- dim_date: ทุกวันของปีที่ครอบทุก date key ----------
    keys = [v for rows in (fgm, fam, fbt, fai) for r in rows for k, v in r.items() if k.endswith("date_key") and v]
    d0, d1 = (datetime.strptime(str(min(keys)), "%Y%m%d"), datetime.strptime(str(max(keys)), "%Y%m%d"))
    dim_date = [dict(zip(["date_key", "full_date", "day_of_week", "month", "month_name", "quarter", "year", "fiscal_year"], r))
                for r in date_rows(d0, d1)]

    tables = {
        "dim_date": dim_date, "dim_department": dep.rows, "dim_user": usr.rows,
        "dim_maintenance_type": mt.rows, "dim_asset_category": cat.rows, "dim_asset_type": typ.rows,
        "dim_location": loc.rows, "dim_asset": ast.rows,
        "fact_borrow_transaction": fbt, "fact_asset_maintenance": fam, "fact_asset_inventory": fai,
        "fact_general_maintenance": [{k: v for k, v in r.items() if k != "_src"} for r in fgm],
    }
    src_count = pd.Series([r["_src"] for r in fgm]).value_counts().to_dict()
    return tables, notes, src_count


WAREHOUSE_ORDER = ["dim_date", "dim_department", "dim_user", "dim_maintenance_type", "dim_asset_category",
                   "dim_asset_type", "dim_location", "dim_asset",
                   "fact_borrow_transaction", "fact_asset_maintenance", "fact_general_maintenance",
                   "fact_asset_inventory"]


def write_load_sql(tables, path, header):
    """full refresh: TRUNCATE ทุกตาราง แล้ว INSERT พร้อม surrogate key ที่ ETL กำหนด"""
    out = ["-- Warehouse full load (MySQL 8) สร้างโดย etl_warehouse.py"] + [f"-- {h}" for h in header] + [
        "SET NAMES utf8mb4;", "SET FOREIGN_KEY_CHECKS = 0;"]
    out += [f"TRUNCATE TABLE `{t}`;" for t in reversed(WAREHOUSE_ORDER)]
    out += ["SET FOREIGN_KEY_CHECKS = 1;", "START TRANSACTION;"]
    for t in WAREHOUSE_ORDER:
        rows = tables[t]
        out += ["", f"-- {t}: {len(rows)} แถว"]
        if not rows:
            continue
        cols = list(rows[0].keys())
        head = f"INSERT INTO `{t}` (" + ", ".join(f"`{c}`" for c in cols) + ") VALUES\n"
        for i in range(0, len(rows), 500):
            out.append(head + ",\n".join("  (" + ", ".join(q(r[c]) for c in cols) + ")" for r in rows[i:i + 500]) + ";")
    out += ["", "COMMIT;", ""]
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(out))


def main(args):
    os.makedirs(args.out_dir, exist_ok=True)
    raw_dir = os.path.join(args.out_dir, "raw")
    wh_dir = os.path.join(args.out_dir, "warehouse")
    os.makedirs(raw_dir, exist_ok=True)
    os.makedirs(wh_dir, exist_ok=True)

    # แหล่งที่ 1: Excel
    excel = run_excel(args.excel, os.path.join(args.out_dir, "repair_requests_clean"))

    # แหล่งที่ 2: API (+ mock สำหรับ endpoint ที่ระบบยังไม่มี)
    api = {}
    for ep in SOURCES:
        api[ep] = fetch(ep, args.mock_dir, raw_dir, offline=args.offline)
        print(f"[api] {ep:<22} {USED[ep]:<28} {len(api[ep]):>5} แถว")

    as_of = datetime.strptime(args.as_of, "%Y-%m-%d") if args.as_of else datetime.now()

    # snapshot ครุภัณฑ์: เก็บรายการ assets ของรอบนี้ไว้ใน <out-dir>/history/ แล้วสร้าง fact_asset_inventory
    # จากทุก snapshot ที่เคยเก็บ (full refresh ของ warehouse จึงไม่ทำประวัติหาย ตราบใดที่โฟลเดอร์ history ยังอยู่)
    history_dir = os.path.join(args.out_dir, "history")
    os.makedirs(history_dir, exist_ok=True)
    if args.snapshot_date:
        snap = datetime.strptime(args.snapshot_date, "%Y-%m-%d")
    elif USED["assets"].startswith("file:"):            # ไฟล์ export: ใช้วันที่ไฟล์ถูกแก้ล่าสุด
        snap = datetime.fromtimestamp(os.path.getmtime(os.path.join(HERE, USED["assets"][len("file:"):])))
    elif USED["assets"] in ("api", "cache"):              # API: วันที่ดึง (ตอน --offline คือวันที่ของ raw ที่ดึงไว้)
        with open(os.path.join(raw_dir, "assets.json"), encoding="utf-8") as f:
            snap = datetime.fromisoformat(json.load(f).get("fetched_at") or as_of.isoformat())
    else:
        snap = as_of
    with open(os.path.join(history_dir, f"assets_{snap:%Y%m%d}.json"), "w", encoding="utf-8") as f:
        json.dump({"snapshot_date": f"{snap:%Y-%m-%d}", "total": len(api["assets"]), "data": api["assets"]}, f, ensure_ascii=False)
    snapshots = load_history(history_dir)
    print(f"[history] snapshot ครุภัณฑ์ {len(snapshots)} ครั้ง: {', '.join(f'{d:%Y-%m-%d}' for d in sorted(snapshots))}")

    tables, notes, src_count = build_warehouse(excel, api, snapshots, as_of)

    mocked = [ep for ep, s in USED.items() if s == "mock"]
    files = [f"{ep} ({s[len('file:'):]})" for ep, s in USED.items() if s.startswith("file:")]
    live = [ep + (" (raw ที่ดึงไว้)" if s == "cache" else "") for ep, s in USED.items() if s in ("api", "cache")]
    fell_back = [ep for ep, s in USED.items() if SOURCES[ep] == "api" and s.startswith("file:")]
    if fell_back:
        notes.insert(0, f"เรียก API ไม่ได้: {', '.join(fell_back)} ใช้ไฟล์ใน repo แทน (snapshot ครุภัณฑ์ลงวันที่ {snap:%Y-%m-%d})")
    header = [f"แหล่งข้อมูล: Excel {os.path.basename(args.excel)} ({src_count.get('excel', 0)} ใบแจ้งซ่อม)",
              f"API {API_BASE}: {', '.join(live) or '-'}",
              f"ไฟล์ข้อมูลจริงใน repo: {', '.join(files) or '-'}",
              f"endpoint จาก MOCK (ไม่ใช่ข้อมูลจริง): {', '.join(mocked)}",
              f"as-of (คำนวณวันเกินกำหนดของรายการยืมที่ยังไม่คืน): {as_of:%Y-%m-%d}",
              "รันซ้ำได้: TRUNCATE ทุกตารางก่อน INSERT (full refresh)",
              f"fact_asset_inventory: snapshot {len(snapshots)} ครั้ง จาก {history_dir} (ห้ามลบโฟลเดอร์นี้ ไม่อย่างนั้นประวัติ snapshot หาย)"] + notes
    write_load_sql(tables, os.path.join(args.out_dir, "warehouse_load.sql"), header)
    for t in WAREHOUSE_ORDER:
        pd.DataFrame(tables[t]).to_csv(os.path.join(wh_dir, f"{t}.csv"), index=False, encoding="utf-8-sig")

    print(f"[warehouse] fact_general_maintenance: {src_count}")
    for t in WAREHOUSE_ORDER:
        print(f"[warehouse] {t:<26} {len(tables[t]):>6} แถว")
    for n in notes:
        print("[note]", n)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description="ETL 2 แหล่ง (Excel ใบแจ้งซ่อม + API ระบบครุภัณฑ์) -> warehouse_load.sql")
    ap.add_argument("excel", nargs="?", help="ไฟล์ .xlsx ใบแจ้งซ่อม (sheet 'การตอบแบบฟอร์ม 1', '... 3', '... 4') "
                                         "default: ไฟล์ .xlsx ในโฟลเดอร์ data/ ของ repo")
    ap.add_argument("-o", "--out-dir", default="output", help="โฟลเดอร์ผลลัพธ์ (default: ./output)")
    ap.add_argument("--mock-dir", default=os.path.join(here, "mock_api"), help="โฟลเดอร์ mock API JSON (default: mock_api ข้างสคริปต์)")
    ap.add_argument("--api-base", default=API_BASE, help="URL ของ API ระบบครุภัณฑ์")
    ap.add_argument("--all-api", action="store_true",
                    help="ดึงทุก endpoint จาก API (ไม่ใช้ mock) เช่น ตอนรันกับ source-api ใน docker compose")
    ap.add_argument("--offline", action="store_true", help="ไม่เรียก API จริง ใช้ raw ที่ดึงไว้ใน <out-dir>/raw/")
    ap.add_argument("--as-of", help="วันที่อ้างอิงคำนวณวันเกินกำหนดของรายการยืมที่ยังไม่คืน YYYY-MM-DD (default: วันนี้)")
    ap.add_argument("--snapshot-date", help="วันที่ของ snapshot ครุภัณฑ์รอบนี้ YYYY-MM-DD "
                                            "(default: วันที่ดึง assets API หรือวันที่แก้ไขไฟล์ assets ล่าสุด)")
    ap.add_argument("--keep-invalid-place", action="store_true",
                    help="เก็บแถวที่สถานที่ใช้ไม่ได้ไว้ (เว้นสถานที่ว่าง) แทนการลบแถว")
    args = ap.parse_args()
    if not args.excel:
        found = sorted(f for f in os.listdir(os.path.join(here, "data")) if f.lower().endswith(".xlsx") and not f.startswith("~$"))
        if not found:
            ap.error("ไม่พบไฟล์ .xlsx ในโฟลเดอร์ data/ ระบุ path ของไฟล์ใบแจ้งซ่อมเอง")
        args.excel = os.path.join(here, "data", found[0])
    API_BASE = args.api_base.rstrip("/") + "/"
    if args.all_api:
        SOURCES.update({ep: "api" for ep in SOURCES})
    if args.keep_invalid_place:
        DROP_INVALID_PLACE = False
    main(args)
